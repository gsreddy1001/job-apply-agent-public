"""Render a tailored markdown resume as an ATS-friendly Word document.

ATS conventions: single column, no tables/text boxes/graphics, black Calibri text,
standard section names in capitals, real bullet lists, each role on one line as
"Title | Company" with dates right-aligned on a tab stop.

Expected markdown (what tailor_job.py produces):
    # Name
    Headline
    contact | contact | contact
    ## Section            (Professional Summary, Skills, Professional Experience, ...)
    ### Title — Company  (then a date line, then "- " bullets)
    - Category: items     (inside Skills)
    - Degree | School | Year   (inside Education)
"""

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Calibri"
BODY_PT = 10.5
MARGIN_TB = Inches(0.6)
MARGIN_LR = Inches(0.7)
TEXT_WIDTH = Inches(8.5) - 2 * MARGIN_LR
BLACK = RGBColor(0, 0, 0)

SECTION_TITLES = {  # Standard ATS section names.
    "summary": "PROFESSIONAL SUMMARY",
    "professional summary": "PROFESSIONAL SUMMARY",
    "skills": "SKILLS",
    "core competencies": "SKILLS",
    "experience": "PROFESSIONAL EXPERIENCE",
    "professional experience": "PROFESSIONAL EXPERIENCE",
    "education": "EDUCATION",
    "certifications": "CERTIFICATIONS",
}


def _font(run, size=BODY_PT, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = BLACK
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    return run


def _spacing(paragraph, before=0, after=0):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = 1.0


def _bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    p_pr.append(borders)


def _right_tab(paragraph):
    paragraph.paragraph_format.tab_stops.add_tab_stop(TEXT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)


def _clean(text: str) -> str:
    text = re.sub(r"\*\*|__", "", text)  # Markdown bold markers.
    # Show URLs as "linkedin.com/in/name": drop the scheme, www, and trailing slash.
    text = re.sub(r"https?://(?:www\.)?(\S+?)/?(?=\s|$)", r"\1", text)
    return re.sub(r"\s{2,}", " ", text).strip()


def _split_role(heading: str) -> tuple[str, str]:
    for sep in (" — ", " – ", " | ", "/ "):
        if sep in heading:
            title, company = heading.rsplit(sep, 1)
            return title.strip(), company.strip()
    return heading.strip(), ""


def _normalize_dates(text: str) -> str:
    text = re.sub(r"\s*[-–—]\s*", " – ", text.strip())
    return re.sub(r"\b([a-z])([a-z]{2})\b", lambda m: m.group(1).upper() + m.group(2), text)


class _Builder:
    def __init__(self):
        self.doc = Document()
        section = self.doc.sections[0]
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = MARGIN_TB
        section.left_margin = section.right_margin = MARGIN_LR
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(BODY_PT)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)

    def para(self, align=None, before=0, after=0):
        p = self.doc.add_paragraph()
        if align is not None:
            p.alignment = align
        _spacing(p, before, after)
        return p

    def name(self, text):
        _font(self.para(WD_ALIGN_PARAGRAPH.CENTER).add_run(text.upper()), size=18, bold=True)

    def centered(self, text, size=BODY_PT, bold=False):
        _font(self.para(WD_ALIGN_PARAGRAPH.CENTER, after=1).add_run(text), size=size, bold=bold)

    def section(self, title):
        p = self.para(before=9, after=3)
        _font(p.add_run(title), size=11.5, bold=True)
        _bottom_border(p)

    def text(self, text):
        p = self.para(after=2)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        _font(p.add_run(text))

    def labeled(self, label, value):
        p = self.para(after=2)
        _font(p.add_run(f"{label}: "), bold=True)
        _font(p.add_run(value))

    def role(self, title, company, dates):
        p = self.para(before=6, after=1)
        _right_tab(p)
        _font(p.add_run(title), bold=True)
        if company:
            _font(p.add_run(f" | {company}"), bold=True)
        if dates:
            _font(p.add_run(f"\t{dates}"))

    def bullet(self, text):
        p = self.doc.add_paragraph(style="List Bullet")
        _spacing(p, after=1)
        fmt = p.paragraph_format
        fmt.left_indent = Inches(0.25)
        fmt.first_line_indent = Inches(-0.18)
        _font(p.add_run(text))

    def education(self, degree, school, year):
        p = self.para(after=1)
        _right_tab(p)
        _font(p.add_run(degree), bold=True)
        if school:
            _font(p.add_run(f", {school}"))
        if year:
            _font(p.add_run(f"\t{year}"))


def build_docx(markdown_text: str, out_path: Path) -> Path:
    b = _Builder()
    lines = [ln.rstrip() for ln in markdown_text.splitlines()]
    section, header_lines, pending_role = None, 0, None

    for raw in lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            b.name(_clean(line[2:]))
            section = "header"
            continue
        if line.startswith("## "):
            key = line[3:].strip().lower()
            section = SECTION_TITLES.get(key, line[3:].strip().upper())
            b.section(section)
            continue
        if line.startswith("### "):
            pending_role = _split_role(_clean(line[4:]))
            continue

        if pending_role is not None:  # The line after a role heading holds its dates.
            if line.startswith("- "):
                b.role(*pending_role, "")
            else:
                b.role(*pending_role, _normalize_dates(line))
                pending_role = None
                continue
            pending_role = None

        text = _clean(line[2:] if line.startswith(("- ", "* ")) else line)
        if section == "header":
            header_lines += 1
            b.centered(text, size=11 if header_lines == 1 else 10, bold=header_lines == 1)
        elif section == "SKILLS" and ":" in text:
            label, value = text.split(":", 1)
            b.labeled(label.strip(), value.strip())
        elif section == "EDUCATION" and "|" in text:
            parts = [p.strip() for p in text.split("|")]
            b.education(parts[0], parts[1] if len(parts) > 1 else "", parts[2] if len(parts) > 2 else "")
        elif section == "CERTIFICATIONS":
            b.text(text)
        elif line.startswith(("- ", "* ")):
            b.bullet(text)
        else:
            b.text(text)

    b.doc.save(out_path)
    return out_path
